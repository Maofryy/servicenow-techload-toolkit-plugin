# /// script
# dependencies = ["python-docx"]
# ///
"""Regenerates ELO-DAT-Template.docx (ELO branding: Century Gothic, ELO blues, logos, confidentiality footer).
Run: uv run make_template.py   (writes next to this file)"""
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

HERE = Path(__file__).parent
FONT = "Century Gothic"
TEXT, NAVY, BLUE, GREY = "1A2433", "003C7D", "0068B3", "5E6B7F"
PALE = "EEF5FC"
CONFIDENTIAL = "© 2026 ELO. Confidential."


def rgb(hex_):
    return RGBColor.from_string(hex_)


def set_font(style, size, color, bold=False, italic=False):
    style.font.name = FONT
    style.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    style.font.size = Pt(size)
    style.font.color.rgb = rgb(color)
    style.font.bold = bold
    style.font.italic = italic


def shade(paragraph, fill):
    pPr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), fill)
    pPr.append(shd)


def left_border(paragraph, color, size=18):
    pPr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    for k, v in (("val", "single"), ("sz", str(size)), ("space", "8"), ("color", color)):
        left.set(qn(f"w:{k}"), v)
    borders.append(left)
    pPr.append(borders)


def add_field(paragraph, instr):
    run = paragraph.add_run()
    for kind, text in (("begin", None), (None, instr), ("separate", None), (None, "1"), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        elif text == instr:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        else:
            el = OxmlElement("w:t")
            el.text = text
        run._r.append(el)
    return run


def define_styles(doc):
    styles = doc.styles
    set_font(styles["Normal"], 10, TEXT)
    styles["Normal"].paragraph_format.space_after = Pt(6)
    styles["Normal"].paragraph_format.line_spacing = 1.15
    specs = {
        "Title": (28, NAVY, True, 0, 6),
        "Subtitle": (14, BLUE, False, 0, 6),
        "Heading 1": (18, NAVY, True, 24, 8),
        "Heading 2": (13, BLUE, True, 16, 6),
        "Heading 3": (11, NAVY, True, 12, 4),
    }
    for name, (size, color, bold, before, after) in specs.items():
        st = styles[name]
        set_font(st, size, color, bold)
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.keep_with_next = True
    for name in ("List Bullet", "List Bullet 2"):
        set_font(styles[name], 10, TEXT)
    callout = styles.add_style("Callout", WD_STYLE_TYPE.PARAGRAPH)
    callout.base_style = styles["Normal"]
    set_font(callout, 9.5, TEXT)
    callout.paragraph_format.left_indent = Cm(0.4)
    callout.paragraph_format.space_before = Pt(6)
    callout.paragraph_format.space_after = Pt(8)
    caption = styles["Caption"]
    set_font(caption, 8.5, GREY, italic=True)
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    code = styles.add_style("Code Inline", WD_STYLE_TYPE.CHARACTER)
    code.font.name = "Menlo"
    code.font.size = Pt(9)
    code.font.color.rgb = rgb(NAVY)


def build_header_footer(section):
    section.different_first_page_header_footer = True
    usable = section.page_width - section.left_margin - section.right_margin
    header = section.header.paragraphs[0]
    header.paragraph_format.tab_stops.add_tab_stop(usable, WD_TAB_ALIGNMENT.RIGHT)
    header.add_run().add_picture(str(HERE / "logo-elo-navy.png"), height=Cm(0.7))
    run = header.add_run("\t{{TITLE}} · {{CLIENT}} · {{VERSION}}")
    run.font.size, run.font.color.rgb = Pt(8), rgb(GREY)
    footer = section.footer.paragraphs[0]
    footer.paragraph_format.tab_stops.add_tab_stop(usable, WD_TAB_ALIGNMENT.RIGHT)
    run = footer.add_run(CONFIDENTIAL + "\tPage ")
    run.font.size, run.font.color.rgb = Pt(8), rgb(GREY)
    for part in (add_field(footer, "PAGE"), footer.add_run(" / "), add_field(footer, "NUMPAGES")):
        part.font.size, part.font.color.rgb = Pt(8), rgb(GREY)


def build_cover(doc):
    doc.add_paragraph().add_run().add_picture(str(HERE / "logo-elo-color.png"), height=Cm(2.2))
    for _ in range(5):
        doc.add_paragraph()
    doc.add_paragraph("Dossier d'Architecture Technique", style="Subtitle")
    doc.add_paragraph("{{TITLE}}", style="Title")
    p = doc.add_paragraph("{{CLIENT}}")
    p.runs[0].font.size, p.runs[0].font.color.rgb = Pt(14), rgb(GREY)
    for _ in range(6):
        doc.add_paragraph()
    meta = doc.add_paragraph()
    meta.add_run("Version {{VERSION}} · {{STATUS}}\n{{DATE}} · {{AUTHOR}}").font.color.rgb = rgb(GREY)
    doc.add_paragraph().add_run().add_picture(str(HERE / "logo-servicenow-dark.png"), height=Cm(0.55))
    q = doc.add_paragraph(CONFIDENTIAL)
    q.runs[0].font.size, q.runs[0].font.color.rgb = Pt(8), rgb(GREY)
    doc.add_page_break()


def build_front_matter(doc):
    doc.add_paragraph("Gestion documentaire", style="Heading 1")
    table = doc.add_table(rows=2, cols=5)
    for cell, text in zip(table.rows[0].cells, ("Version", "Date", "Auteur", "Statut", "Modifications")):
        cell.text = text
    for cell, text in zip(table.rows[1].cells, ("{{VERSION}}", "{{DATE}}", "{{AUTHOR}}", "{{STATUS}}", "{{CHANGES}}")):
        cell.text = text
    doc.add_paragraph()
    doc.add_paragraph("Table des matières", style="Heading 1")
    add_field(doc.add_paragraph(), 'TOC \\o "1-3" \\h \\z \\u')
    doc.add_page_break()
    doc.add_paragraph("{{BODY}}")


def main():
    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.left_margin = section.right_margin = Cm(2.2)
    section.top_margin, section.bottom_margin = Cm(2.4), Cm(2.2)
    define_styles(doc)
    build_header_footer(section)
    build_cover(doc)
    build_front_matter(doc)
    settings = doc.settings.element
    update = OxmlElement("w:updateFields")
    update.set(qn("w:val"), "true")
    settings.append(update)
    out = HERE / "ELO-DAT-Template.docx"
    doc.save(out)
    print(out)


if __name__ == "__main__":
    main()
